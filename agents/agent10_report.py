"""
Agent 10 — Report Generator
Assembles the full contract review report and exports to PDF, DOCX, and JSON.
"""
import json
import io
from datetime import datetime
from typing import Dict, List, Any

# ── PDF ──────────────────────────────────────────────────────────────────────
def generate_pdf(state: Dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        HRFlowable, PageBreak,
    )
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            rightMargin=20*mm, leftMargin=20*mm,
                            topMargin=20*mm, bottomMargin=20*mm)

    NAVY  = colors.HexColor("#1A2B5E")
    BLUE  = colors.HexColor("#2E75B6")
    RED   = colors.HexColor("#C00000")
    AMBER = colors.HexColor("#E07000")
    GREEN = colors.HexColor("#2E8B57")
    LIGHT = colors.HexColor("#EBF3FB")

    styles = getSampleStyleSheet()
    title_s  = ParagraphStyle("title",  fontName="Helvetica-Bold", fontSize=22, textColor=NAVY,  alignment=TA_CENTER, spaceAfter=6)
    sub_s    = ParagraphStyle("sub",    fontName="Helvetica",      fontSize=12, textColor=BLUE,  alignment=TA_CENTER, spaceAfter=4)
    h1_s     = ParagraphStyle("h1",     fontName="Helvetica-Bold", fontSize=14, textColor=NAVY,  spaceBefore=12, spaceAfter=4)
    h2_s     = ParagraphStyle("h2",     fontName="Helvetica-Bold", fontSize=11, textColor=BLUE,  spaceBefore=8,  spaceAfter=3)
    body_s   = ParagraphStyle("body",   fontName="Helvetica",      fontSize=9,  textColor=colors.HexColor("#2C2C2C"), spaceAfter=4, leading=13)
    small_s  = ParagraphStyle("small",  fontName="Helvetica",      fontSize=8,  textColor=colors.grey, spaceAfter=2)
    risk_h   = ParagraphStyle("riskh",  fontName="Helvetica-Bold", fontSize=10, textColor=RED,   spaceBefore=6, spaceAfter=2)
    green_s  = ParagraphStyle("green",  fontName="Helvetica-Bold", fontSize=9,  textColor=GREEN)

    def hr(): return HRFlowable(width="100%", thickness=1, color=BLUE, spaceAfter=6)
    def risk_color(lvl):
        return {
            "HIGH": RED, "MEDIUM": AMBER, "LOW": GREEN, "NONE": colors.grey
        }.get(lvl, colors.grey)

    meta   = state.get("extracted_metadata", {})
    clauses= state.get("clauses", [])
    score  = state.get("risk_score", 0)
    bd     = state.get("risk_breakdown", {})
    recs   = state.get("recommendations", [])
    fname  = state.get("filename", "Contract")

    story = []

    # Cover
    story.append(Spacer(1, 20*mm))
    story.append(Paragraph("LexaGuard AI", title_s))
    story.append(Paragraph("Legal Contract Intelligence Report", sub_s))
    story.append(Spacer(1, 4*mm))
    story.append(hr())
    story.append(Paragraph(f"Document: {fname}", body_s))
    story.append(Paragraph(f"Contract Type: {meta.get('contract_type','Unknown')}", body_s))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%d %B %Y, %H:%M')}", small_s))
    story.append(Spacer(1, 6*mm))

    # Risk score box
    score_data = [[
        Paragraph(f"Overall Risk Score", sub_s),
        Paragraph(f"{score}/100", ParagraphStyle("sc", fontName="Helvetica-Bold", fontSize=28,
                  textColor=RED if score>60 else AMBER if score>30 else GREEN, alignment=TA_CENTER)),
    ]]
    st = TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), LIGHT),
        ("BOX", (0,0), (-1,-1), 1, BLUE),
        ("ALIGN", (0,0), (-1,-1), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,0), (-1,-1), [LIGHT]),
    ])
    story.append(Table(score_data, colWidths=[90*mm, 70*mm], style=st))
    story.append(Spacer(1, 4*mm))

    # Risk breakdown
    bd_data = [["Dimension", "Rating"]] + [
        [k.title(), v] for k, v in bd.items() if k in ("financial","legal","operational")
    ]
    bd_style = TableStyle([
        ("BACKGROUND", (0,0), (-1,0), NAVY),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 9),
        ("ROWBACKGROUNDS", (1,0), (-1,-1), [LIGHT, colors.white]),
        ("BOX", (0,0), (-1,-1), 0.5, BLUE),
        ("GRID", (0,0), (-1,-1), 0.3, colors.lightgrey),
    ])
    story.append(Table(bd_data, colWidths=[80*mm, 80*mm], style=bd_style))
    story.append(PageBreak())

    # Executive summary
    story.append(Paragraph("Executive Summary", h1_s))
    story.append(hr())
    story.append(Paragraph(state.get("executive_summary", ""), body_s))
    story.append(Spacer(1, 4*mm))

    # Key metadata
    story.append(Paragraph("Contract Details", h1_s))
    story.append(hr())
    meta_rows = [
        ["Parties", f"{meta.get('parties',{}).get('party_a','?')} ↔ {meta.get('parties',{}).get('party_b','?')}"],
        ["Effective Date", meta.get("effective_date","Unknown")],
        ["Expiry Date", meta.get("expiry_date","Unknown")],
        ["Contract Value", meta.get("contract_value","Unknown")],
        ["Payment Terms", meta.get("payment_terms","Unknown")],
        ["Termination Notice", meta.get("termination_notice","Unknown")],
        ["Governing Law", meta.get("governing_law","Unknown")],
    ]
    meta_tbl = Table(meta_rows, colWidths=[60*mm, 110*mm])
    meta_tbl.setStyle(TableStyle([
        ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 9),
        ("ROWBACKGROUNDS", (0,0), (-1,-1), [LIGHT, colors.white]),
        ("BOX", (0,0), (-1,-1), 0.5, BLUE),
        ("GRID", (0,0), (-1,-1), 0.3, colors.lightgrey),
    ]))
    story.append(meta_tbl)
    story.append(Spacer(1, 4*mm))

    # High risk clauses
    high = [c for c in clauses if c.get("risk_level") == "HIGH"]
    if high:
        story.append(Paragraph("⚠ High Risk Clauses", h1_s))
        story.append(hr())
        for c in high:
            story.append(Paragraph(f"{c['id']} — {c.get('category','')} | Issue: {c.get('risk_issue','')}", risk_h))
            story.append(Paragraph(f"<b>Original:</b> {c['text'][:300]}...", body_s))
            if c.get("legal_reasoning"):
                story.append(Paragraph(f"<b>Legal Analysis:</b> {c['legal_reasoning']}", body_s))
            if c.get("redline_suggestion"):
                story.append(Paragraph(f"<b>Suggested Redline:</b> {c['redline_suggestion']}", body_s))
            if c.get("plain_english"):
                story.append(Paragraph(f"<b>Plain English:</b> {c['plain_english']}", body_s))
            story.append(Spacer(1, 3*mm))

    # Recommendations
    if recs:
        story.append(PageBreak())
        story.append(Paragraph("Recommendations", h1_s))
        story.append(hr())
        for r in recs:
            story.append(Paragraph(f"• {r}", body_s))

    # Plain summary
    story.append(PageBreak())
    story.append(Paragraph("Plain Language Summary", h1_s))
    story.append(hr())
    story.append(Paragraph(state.get("plain_summary", ""), body_s))

    doc.build(story)
    return buf.getvalue()


# ── DOCX ─────────────────────────────────────────────────────────────────────
def generate_docx(state: Dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    NAVY  = RGBColor(0x1A, 0x2B, 0x5E)
    BLUE  = RGBColor(0x2E, 0x75, 0xB6)
    RED   = RGBColor(0xC0, 0x00, 0x00)
    AMBER = RGBColor(0xE0, 0x70, 0x00)
    GREEN = RGBColor(0x2E, 0x8B, 0x57)

    def add_h(text, color=NAVY, size=16, bold=True):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = color
        return p

    def add_p(text, color=RGBColor(0x2C,0x2C,0x2C), size=10):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(size); r.font.color.rgb = color
        return p

    meta    = state.get("extracted_metadata", {})
    clauses = state.get("clauses", [])
    score   = state.get("risk_score", 0)

    add_h("LexaGuard AI — Contract Review Report", NAVY, 18)
    add_p(f"File: {state.get('filename','')} | Generated: {datetime.now().strftime('%d %B %Y')}", BLUE, 9)
    doc.add_paragraph()

    add_h(f"Overall Risk Score: {score}/100", RED if score>60 else AMBER if score>30 else GREEN, 14)
    doc.add_paragraph()

    add_h("Executive Summary", NAVY, 13)
    add_p(state.get("executive_summary",""))
    doc.add_paragraph()

    add_h("High Risk Clauses", RED, 13)
    for c in [x for x in clauses if x.get("risk_level") == "HIGH"]:
        add_h(f"{c['id']} — {c.get('category','')} — {c.get('risk_issue','')}", RED, 11)
        add_p(f"Original: {c['text'][:400]}")
        if c.get("redline_suggestion"):
            p = doc.add_paragraph()
            r = p.add_run("Suggested Redline: ")
            r.bold = True; r.font.color.rgb = GREEN
            p.add_run(c["redline_suggestion"])
        if c.get("negotiation_tip"):
            add_p(f"Tip: {c['negotiation_tip']}", AMBER)
        doc.add_paragraph()

    add_h("Recommendations", NAVY, 13)
    for rec in state.get("recommendations", []):
        doc.add_paragraph(f"• {rec}", style="List Bullet")

    add_h("Plain Language Summary", NAVY, 13)
    add_p(state.get("plain_summary",""))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── JSON ──────────────────────────────────────────────────────────────────────
def generate_json(state: Dict) -> str:
    export = {
        "filename": state.get("filename"),
        "generated_at": datetime.now().isoformat(),
        "risk_score": state.get("risk_score"),
        "risk_breakdown": state.get("risk_breakdown"),
        "executive_summary": state.get("executive_summary"),
        "plain_summary": state.get("plain_summary"),
        "recommendations": state.get("recommendations"),
        "extracted_metadata": state.get("extracted_metadata"),
        "clauses": [
            {k: v for k, v in c.items() if k != "text"}
            for c in state.get("clauses", [])
        ],
    }
    return json.dumps(export, indent=2, default=str)


# ── Executive Summary & Recommendations (LLM) ────────────────────────────────
def generate_summaries(state: Dict, client) -> Dict:
    from utils.llm_client import chat
    import re

    clauses  = state.get("clauses", [])
    score    = state.get("risk_score", 0)
    bd       = state.get("risk_breakdown", {})
    meta     = state.get("extracted_metadata", {})
    high_cls = [c for c in clauses if c.get("risk_level") == "HIGH"]
    med_cls  = [c for c in clauses if c.get("risk_level") == "MEDIUM"]

    context = (
        f"Contract: {meta.get('contract_type','Unknown')}\n"
        f"Parties: {meta.get('parties',{})}\n"
        f"Risk Score: {score}/100\n"
        f"High Risk Clauses ({len(high_cls)}): "
        + "; ".join(f"{c['id']} {c.get('risk_issue','')}" for c in high_cls[:5])
        + f"\nMedium Risk ({len(med_cls)} clauses)\n"
        f"Financial Risk: {bd.get('financial')}, Legal: {bd.get('legal')}, Operational: {bd.get('operational')}"
    )

    exec_prompt = f"""You are a senior legal partner writing an executive summary for a client.

Contract analysis:
{context}

Write a 3-4 sentence executive summary covering:
1. Overall risk assessment
2. Most critical issues
3. Recommended immediate actions

Be direct and professional. No bullet points."""

    rec_prompt = f"""Based on this contract analysis:
{context}

Provide exactly 6 specific, actionable recommendations to reduce legal risk.
Return ONLY a JSON array of strings. Each recommendation under 25 words."""

    plain_prompt = f"""Write a plain-language summary of this contract for a non-lawyer.

{context}

High risk issues: {[c.get('plain_english','') for c in high_cls[:3]]}

Write 4-5 sentences in simple English. No jargon. Focus on what the signing party needs to know."""

    try:
        exec_summary = chat(client, [{"role":"user","content": exec_prompt}], temperature=0.3)
    except Exception:
        exec_summary = f"This contract has an overall risk score of {score}/100. There are {len(high_cls)} high-risk clauses requiring immediate attention."

    try:
        rec_raw = chat(client, [{"role":"user","content": rec_prompt}], temperature=0.2)
        rec_raw = re.sub(r"```[a-z]*\n?", "", rec_raw).replace("```","").strip()
        recs = json.loads(rec_raw)
        if not isinstance(recs, list):
            raise ValueError
    except Exception:
        recs = [
            f"Review and redline the {c.get('category','')} clause ({c['id']}): {c.get('risk_issue','')}"
            for c in high_cls[:4]
        ] + ["Engage legal counsel before signing.", "Request a 30-day review period."]

    try:
        plain = chat(client, [{"role":"user","content": plain_prompt}], temperature=0.4)
    except Exception:
        plain = f"This is a {meta.get('contract_type','contract')} with {len(clauses)} clauses. It has {len(high_cls)} serious issues that need attention before signing."

    state["executive_summary"] = exec_summary
    state["recommendations"]   = recs[:8]
    state["plain_summary"]     = plain
    state["report_ready"]      = True
    return state
